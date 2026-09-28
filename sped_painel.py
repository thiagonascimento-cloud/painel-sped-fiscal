"""Gera a planilha de Compras e Vendas a partir de arquivos SPED Fiscal (EFD ICMS/IPI).

Uso:
    python sped_painel.py                      -> abre janela para escolher os .txt
    python sped_painel.py <pasta ou arquivos>  -> processa todos os *.txt SPED informados
    python sped_painel.py <...> -o saida.xlsx  -> define o arquivo de saída

Abas geradas (mesmo formato do botão "Baixar Excel" do Painel SPED.html):
    Resumo                   dados da empresa e valores consolidados (só o que tem valor)
    Notas fiscais            uma linha por nota x CFOP
    Produtos                 itens das notas (registro C170)
    Clientes e fornecedores  totais por participante
Colunas e linhas sem valor não são geradas; zeros ficam em branco.
"""
import argparse
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

CFOP = {
    1101: ('Compra para industrialização', 'Compra para revenda'),
    1102: ('Compra para comercialização', 'Compra para revenda'),
    2102: ('Compra para comercialização (interestadual)', 'Compra para revenda'),
    1403: ('Compra para comercialização em operação com mercadoria sujeita a ST', 'Compra para revenda'),
    2403: ('Compra para comercialização em operação com mercadoria sujeita a ST (interestadual)', 'Compra para revenda'),
    1551: ('Compra de bem para o ativo imobilizado', 'Compra uso/consumo/ativo'),
    2551: ('Compra de bem para o ativo imobilizado (interestadual)', 'Compra uso/consumo/ativo'),
    1556: ('Compra de material para uso ou consumo', 'Compra uso/consumo/ativo'),
    2556: ('Compra de material para uso ou consumo (interestadual)', 'Compra uso/consumo/ativo'),
    1406: ('Compra de bem para o ativo imobilizado sujeito a ST', 'Compra uso/consumo/ativo'),
    2406: ('Compra de bem para o ativo imobilizado sujeito a ST (interestadual)', 'Compra uso/consumo/ativo'),
    1407: ('Compra de mercadoria para uso ou consumo sujeita a ST', 'Compra uso/consumo/ativo'),
    2407: ('Compra de mercadoria para uso ou consumo sujeita a ST (interestadual)', 'Compra uso/consumo/ativo'),
    1201: ('Devolução de venda de produção do estabelecimento', 'Devolução de venda'),
    1202: ('Devolução de venda de mercadoria adquirida ou recebida de terceiros', 'Devolução de venda'),
    2202: ('Devolução de venda de mercadoria adquirida ou recebida de terceiros (interestadual)', 'Devolução de venda'),
    1410: ('Devolução de venda de produção do estabelecimento em operação com ST', 'Devolução de venda'),
    1411: ('Devolução de venda de mercadoria adquirida ou recebida de terceiros em operação com ST', 'Devolução de venda'),
    2411: ('Devolução de venda de mercadoria em operação com ST (interestadual)', 'Devolução de venda'),
    1152: ('Transferência para comercialização', 'Transferência'),
    2152: ('Transferência para comercialização (interestadual)', 'Transferência'),
    1409: ('Transferência para comercialização em operação com ST', 'Transferência'),
    2409: ('Transferência para comercialização em operação com ST (interestadual)', 'Transferência'),
    1552: ('Transferência de bem do ativo imobilizado', 'Transferência'),
    1557: ('Transferência de material de uso ou consumo', 'Transferência'),
    1253: ('Compra de energia elétrica por estabelecimento comercial', 'Outras'),
    1303: ('Aquisição de serviço de comunicação por estabelecimento comercial', 'Outras'),
    1353: ('Aquisição de serviço de transporte por estabelecimento comercial', 'Serviços/Fretes'),
    2353: ('Aquisição de serviço de transporte por estabelecimento comercial (interestadual)', 'Serviços/Fretes'),
    1933: ('Aquisição de serviço tributado pelo ISSQN', 'Serviços/Fretes'),
    1904: ('Retorno de remessa para venda fora do estabelecimento', 'Remessa/Retorno'),
    1910: ('Entrada de bonificação, doação ou brinde', 'Bonificação/Brinde'),
    2910: ('Entrada de bonificação, doação ou brinde (interestadual)', 'Bonificação/Brinde'),
    1912: ('Entrada de mercadoria ou bem recebido para demonstração ou mostruário', 'Remessa/Retorno'),
    1913: ('Retorno de mercadoria ou bem remetido para demonstração, mostruário ou treinamento', 'Remessa/Retorno'),
    1914: ('Retorno de mercadoria ou bem remetido para exposição ou feira', 'Remessa/Retorno'),
    2914: ('Retorno de mercadoria ou bem remetido para exposição ou feira (interestadual)', 'Remessa/Retorno'),
    1915: ('Entrada de mercadoria ou bem recebido para conserto ou reparo', 'Remessa/Retorno'),
    1916: ('Retorno de mercadoria ou bem remetido para conserto ou reparo', 'Remessa/Retorno'),
    1949: ('Outra entrada de mercadoria ou prestação de serviço não especificada', 'Outras'),
    2949: ('Outra entrada de mercadoria ou prestação de serviço não especificada (interestadual)', 'Outras'),
    5101: ('Venda de produção do estabelecimento', 'Venda'),
    5102: ('Venda de mercadoria adquirida ou recebida de terceiros', 'Venda'),
    6102: ('Venda de mercadoria adquirida ou recebida de terceiros (interestadual)', 'Venda'),
    6108: ('Venda de mercadoria adquirida ou recebida de terceiros, destinada a não contribuinte', 'Venda'),
    5403: ('Venda de mercadoria adquirida ou recebida de terceiros em operação com ST, na condição de contribuinte substituto', 'Venda'),
    6403: ('Venda de mercadoria em operação com ST, na condição de contribuinte substituto (interestadual)', 'Venda'),
    5405: ('Venda de mercadoria adquirida ou recebida de terceiros em operação com ST, na condição de contribuinte substituído', 'Venda'),
    6404: ('Venda de mercadoria sujeita a ST, cujo imposto já tenha sido retido anteriormente', 'Venda'),
    5551: ('Venda de bem do ativo imobilizado', 'Venda de ativo'),
    5933: ('Prestação de serviço tributado pelo ISSQN', 'Serviços/Fretes'),
    5202: ('Devolução de compra para comercialização', 'Devolução de compra'),
    6202: ('Devolução de compra para comercialização (interestadual)', 'Devolução de compra'),
    5411: ('Devolução de compra para comercialização em operação com ST', 'Devolução de compra'),
    6411: ('Devolução de compra para comercialização em operação com ST (interestadual)', 'Devolução de compra'),
    5556: ('Devolução de compra de material de uso ou consumo', 'Devolução de compra'),
    5152: ('Transferência de mercadoria adquirida ou recebida de terceiros', 'Transferência'),
    6152: ('Transferência de mercadoria adquirida ou recebida de terceiros (interestadual)', 'Transferência'),
    5409: ('Transferência de mercadoria adquirida ou recebida de terceiros em operação com ST', 'Transferência'),
    6409: ('Transferência de mercadoria em operação com ST (interestadual)', 'Transferência'),
    5552: ('Transferência de bem do ativo imobilizado', 'Transferência'),
    5557: ('Transferência de material de uso ou consumo', 'Transferência'),
    5904: ('Remessa para venda fora do estabelecimento', 'Remessa/Retorno'),
    5910: ('Remessa em bonificação, doação ou brinde', 'Bonificação/Brinde'),
    6910: ('Remessa em bonificação, doação ou brinde (interestadual)', 'Bonificação/Brinde'),
    5912: ('Remessa de mercadoria ou bem para demonstração, mostruário ou treinamento', 'Remessa/Retorno'),
    6912: ('Remessa de mercadoria ou bem para demonstração, mostruário ou treinamento (interestadual)', 'Remessa/Retorno'),
    5913: ('Retorno de mercadoria ou bem recebido para demonstração ou mostruário', 'Remessa/Retorno'),
    5914: ('Remessa de mercadoria ou bem para exposição ou feira', 'Remessa/Retorno'),
    6914: ('Remessa de mercadoria ou bem para exposição ou feira (interestadual)', 'Remessa/Retorno'),
    5915: ('Remessa de mercadoria ou bem para conserto ou reparo', 'Remessa/Retorno'),
    5916: ('Retorno de mercadoria ou bem recebido para conserto ou reparo', 'Remessa/Retorno'),
    5927: ('Baixa de estoque decorrente de perda, roubo ou deterioração', 'Outras'),
    5929: ('Documento fiscal relativo a operação também registrada em ECF', 'Outras'),
    5949: ('Outra saída de mercadoria ou prestação de serviço não especificado', 'Outras'),
    6949: ('Outra saída de mercadoria ou prestação de serviço não especificado (interestadual)', 'Outras'),
}


def regra_grupo(c):
    """Natureza pelo grupo do CFOP (para códigos fora da tabela oficial)."""
    g, ent = c % 1000, c // 1000 <= 3
    if g in (910, 911):
        return 'Bonificação/Brinde'
    if 200 <= g <= 219 or 410 <= g <= 413 or g in (553, 555, 918, 919) or 660 <= g <= 662:
        return 'Devolução de venda' if ent else 'Devolução de compra'
    if 150 <= g <= 159 or g in (408, 409, 552, 557, 658, 659):
        return 'Transferência'
    if 350 <= g <= 360 or g in (932, 933):
        return 'Serviços/Fretes'
    if (g in (414, 415, 554, 657, 663, 664, 665, 666) or 501 <= g <= 505
            or (900 <= g <= 949 and g not in (922, 926, 927, 928, 929, 931, 949))):
        return 'Remessa/Retorno'
    if ent:
        if g in (124, 125):
            return 'Outras'
        if 100 <= g <= 149 or g in (401, 403, 651, 652):
            return 'Compra para revenda'
        return 'Compra uso/consumo/ativo' if g in (406, 407, 551, 556, 653) else 'Outras'
    if 100 <= g <= 149 or 401 <= g <= 405 or 651 <= g <= 656 or g == 667 or 250 <= g <= 259 or 300 <= g <= 309:
        return 'Venda'
    return 'Venda de ativo' if g == 551 else 'Outras'


def carregar_tabela_cfop():
    """Tabela oficial (IT 2023.002 v2.10, 619 CFOPs) em cfop_oficial.json, ao lado deste script.
    Retorna {cfop: (descrição, natureza, classificação, nota explicativa)}."""
    arq = Path(__file__).with_name('cfop_oficial.json')
    if arq.exists():
        import json
        return {c: (t, n, o, d) for c, t, n, o, d in json.loads(arq.read_text(encoding='utf-8'))}
    return {c: (d, n, 'revisada', '') for c, (d, n) in CFOP.items()}


CFOP_TAB = carregar_tabela_cfop()


def cfop_info(c):
    if c in CFOP_TAB:
        return CFOP_TAB[c]
    return ('CFOP não consta na tabela oficial (IT 2023.002 v2.10) – verificar', regra_grupo(c), 'automática', '')
UF = {'11': 'RO', '12': 'AC', '13': 'AM', '14': 'RR', '15': 'PA', '16': 'AP', '17': 'TO', '21': 'MA', '22': 'PI',
      '23': 'CE', '24': 'RN', '25': 'PB', '26': 'PE', '27': 'AL', '28': 'SE', '29': 'BA', '31': 'MG', '32': 'ES',
      '33': 'RJ', '35': 'SP', '41': 'PR', '42': 'SC', '43': 'RS', '50': 'MS', '51': 'MT', '52': 'GO', '53': 'DF'}
SIT = {'00': 'Regular', '01': 'Regular extemporâneo', '02': 'Cancelado', '03': 'Cancelado extemporâneo',
       '04': 'Denegado', '05': 'Numeração inutilizada', '06': 'Complementar', '07': 'Complementar extemporâneo',
       '08': 'Regime especial'}
DIAS = ['seg', 'ter', 'qua', 'qui', 'sex', 'sáb', 'dom']

NAVY, LIGHT, GREEN, RED = '1F3A4D', 'E8EFF7', '2E8A4A', 'B03C3C'
F_BASE = Font(name='Arial', size=10)
F_HDR = Font(name='Arial', size=10, bold=True, color='FFFFFF')
FILL_HDR = PatternFill('solid', fgColor=NAVY)
FILL_LIGHT = PatternFill('solid', fgColor=LIGHT)
THIN = Side(style='thin', color='9FB1BF')
HAIR = Side(style='thin', color='E3E8EC')
MD = "'Movimento Diário'!"


# ---------------------------------------------------------------- leitura
def num(s):
    return float(s.replace(',', '.')) if s else 0.0


def dt(s):
    try:
        return datetime.strptime(s, '%d%m%Y').date() if s and len(s) == 8 else None
    except ValueError:
        return None


def ler_texto(path):
    raw = Path(path).read_bytes()
    try:
        return raw.decode('utf-8')
    except UnicodeDecodeError:
        return raw.decode('cp1252', errors='replace')


def parse_sped(texto, origem=''):
    est = {'rows': [], 'itens': [], 'arquivo': origem}
    part, prod, cur = {}, {}, None

    def item(cst, cfop, vl, bc, icms, bcst, st, ipi):
        if cur is None:
            return
        est['rows'].append({**cur, 'cfop': int(cfop) if cfop else None, 'cst': cst, 'vl': num(vl), 'bc': num(bc),
                            'icms': num(icms), 'bcst': num(bcst), 'st': num(st), 'ipi': num(ipi),
                            'first': 0 if cur['_has'] else 1})
        cur['_has'] = True

    def hdr(oper, emit, cp, mod, sit, ser, nro, chv, d_doc, d_es, vl, reg):
        nonlocal cur
        p = part.get(cp)
        cur = {'es': 'Entrada' if oper == '0' else 'Saída', 'emit': 'Própria' if emit == '0' else 'Terceiros',
               'part': p['nome'] if p else 'Consumidor / não identificado', 'doc': p['doc'] if p else '',
               'uf': p['uf'] if p else '', 'mod': mod, 'sit': SIT.get(sit, sit), 'ser': ser, 'nro': nro,
               'chv': chv or '', 'd_doc': dt(d_doc), 'data': dt(d_es) or dt(d_doc), 'reg': reg, '_has': False}
        if sit in ('02', '03', '04', '05'):
            item('', '', '0', '0', '0', '0', '0', '0')

    for ln in texto.splitlines():
        if ln.startswith('|9999|'):
            break
        if not ln.startswith('|'):
            continue
        if '&' in ln:
            ln = ln.replace('&AMP;', '&').replace('&amp;', '&')
        f = ln.split('|') + [''] * 30
        r = f[1]
        if r == '0000':
            est.update(ini=dt(f[4]), fim=dt(f[5]), empresa=f[6], cnpj=f[7], uf=f[9], ie=f[10])
        elif r == '0150':
            part[f[2]] = {'nome': f[3], 'doc': f[5] or f[6], 'uf': UF.get(f[8][:2], '')}
        elif r == '0200':
            prod[f[2]] = {'desc': f[3], 'un': f[6], 'ncm': f[8]}
        elif r == 'C170' and cur is not None:
            p = prod.get(f[3], {})
            est['itens'].append({'data': cur['data'], 'es': cur['es'], 'nro': cur['nro'], 'part': cur['part'],
                                 'item': int(f[2]) if f[2].isdigit() else 0, 'cod': f[3],
                                 'desc': p.get('desc') or f[4] or '(produto sem cadastro no 0200)', 'compl': f[4],
                                 'ncm': p.get('ncm', ''), 'qtd': num(f[5]), 'un': f[6] or p.get('un', ''),
                                 'vl': num(f[7]), 'dsc': num(f[8]), 'cfop': int(f[11]) if f[11].isdigit() else None,
                                 'bc': num(f[13]), 'icms': num(f[15]), 'bcst': num(f[16]), 'st': num(f[18]),
                                 'ipi': num(f[24])})
        elif r == 'C100':
            hdr(f[2], f[3], f[4], f[5], f[6], f[7], f[8], f[9], f[10], f[11], f[12], 'C100')
        elif r == 'C190':
            item(f[2], f[3], f[5], f[6], f[7], f[8], f[9], f[11])
        elif r == 'C500':
            hdr(f[2], f[3], f[4], f[5], f[6], f[7], f[10], '', f[11], f[12], f[13], 'C500')
        elif r == 'C590':
            item(f[2], f[3], f[5], f[6], f[7], f[8], f[9], '0')
        elif r == 'D100':
            hdr(f[2], f[3], f[4], f[5], f[6], f[7], f[9], f[10], f[11], f[12], f[15], 'D100')
        elif r in ('D190', 'D590'):
            item(f[2], f[3], f[5], f[6], f[7], '0', '0', '0')
        elif r == 'D500':
            hdr(f[2], f[3], f[4], f[5], f[6], f[7], f[9], '', f[10], f[11], f[12], 'D500')
    if 'cnpj' not in est:
        raise ValueError(f'{origem}: registro 0000 não encontrado (não parece um SPED Fiscal EFD ICMS/IPI).')
    return est


def rotulos(ests):
    bases = {e['cnpj'][:8] for e in ests}
    periodos = {e['ini'] for e in ests}
    for e in ests:
        o, dv = e['cnpj'][8:12], e['cnpj'][12:]
        lb = ('Matriz ' if o == '0001' else 'Filial ') + f'{o}-{dv}'
        if len(bases) > 1:
            lb = ' '.join(e['empresa'].split()[:3]) + ' · ' + lb
        if len(periodos) > 1:
            lb += ' · ' + e['ini'].strftime('%m/%Y')
        e['label'] = lb


# ---------------------------------------------------------------- helpers de estilo
def hdr_row(ws, row, c1, c2):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row, c)
        cell.font, cell.fill = F_HDR, FILL_HDR
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def total_row(ws, row, c1, c2):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row, c)
        cell.font = Font(name='Arial', size=10, bold=True)
        cell.fill = FILL_LIGHT
        cell.border = Border(top=THIN)


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------------------------------------------------------------- planilha (formato simplificado)
MONEY = '#,##0.00;-#,##0.00;""'
F_TIT = Font(name='Arial', size=14, bold=True, color=NAVY)
F_SEC = Font(name='Arial', size=11, bold=True, color=NAVY)
F_NOTA = Font(name='Arial', size=9, italic=True, color='595959')
F_BOLD = Font(name='Arial', size=10, bold=True)


def nz(v):
    return v is not None and abs(v) >= 0.005


def fmt_doc(d):
    if d and len(d) == 14:
        return f'{d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:]}'
    if d and len(d) == 11:
        return f'{d[:3]}.{d[3:6]}.{d[6:9]}-{d[9:]}'
    return d or ''


def tabela(ws, r0, cols, dados, total=True):
    """Escreve uma tabela a partir da linha r0 (cabeçalho). Colunas numéricas sem nenhum valor são omitidas
    e zeros ficam em branco. Retorna ({chave: letra da coluna}, linha do total ou None)."""
    vis = [c for c in cols if not c.get('num') or c.get('fx') or any(nz(c['f'](d)) for d in dados)]
    letra = {c['k']: get_column_letter(i) for i, c in enumerate(vis, 1)}
    for j, c in enumerate(vis, 1):
        ws.cell(r0, j, c['t'])
        ws.column_dimensions[get_column_letter(j)].width = c['w']
    hdr_row(ws, r0, 1, len(vis))
    ws.row_dimensions[r0].height = 28
    for i, d in enumerate(dados, r0 + 1):
        for j, c in enumerate(vis, 1):
            if c.get('fx'):
                v = c['fx'](i, letra)
            else:
                v = c['f'](d)
                if c.get('num') and not nz(v):
                    v = None
            cell = ws.cell(i, j, v)
            cell.font = F_BASE
            if c.get('date'):
                cell.number_format = 'dd/mm/yyyy'
            elif c.get('z'):
                cell.number_format = c['z']
            elif c.get('num'):
                cell.number_format = MONEY
    fim = r0 + len(dados)
    ws.auto_filter.ref = f'A{r0}:{get_column_letter(len(vis))}{max(fim, r0 + 1)}'
    tot = None
    if total and dados:
        tot = fim + 1
        ws.cell(tot, 1, 'Total')
        for j, c in enumerate(vis, 1):
            if c.get('num') and c.get('tot', True):
                L = get_column_letter(j)
                cell = ws.cell(tot, j, f'=SUBTOTAL(9,{L}{r0 + 1}:{L}{fim})')
                cell.number_format = c.get('z') or MONEY
        total_row(ws, tot, 1, len(vis))
    return letra, tot


def gerar(ests, saida):
    rotulos(ests)
    rows, itens = [], []
    for e in ests:
        for r in e['rows']:
            d, n, _, _ = cfop_info(r['cfop']) if r['cfop'] else ('', '', '', '')
            rows.append({**r, 'filial': e['label'], 'desc': d, 'nat': n})
        for it in e['itens']:
            itens.append({**it, 'filial': e['label'], 'nat': cfop_info(it['cfop'])[1] if it['cfop'] else ''})
    chave_nf = lambda r: int(r['nro']) if str(r['nro']).isdigit() else 0
    rows.sort(key=lambda r: (r['data'] or date.min, r['filial'], r['es'], chave_nf(r)))
    itens.sort(key=lambda r: (r['data'] or date.min, chave_nf(r), r['item']))
    multi = sum(1 for e in ests if e['rows']) > 1
    empresa = ' / '.join(dict.fromkeys(e['empresa'] for e in ests))
    ini, fim = min(e['ini'] for e in ests), max(e['fim'] for e in ests)
    per = ini.strftime('%m/%Y')

    wb = Workbook()
    wsR = wb.active
    wsR.title = 'Resumo'
    wsN = wb.create_sheet('Notas fiscais')
    wsP = wb.create_sheet('Produtos')
    wsC = wb.create_sheet('Clientes e fornecedores')
    for ws in (wsR, wsN, wsP, wsC):
        ws.sheet_view.showGridLines = False

    # ---- Notas fiscais (uma linha por nota x CFOP; sem as notas zeradas)
    notas = [r for r in rows if nz(r['vl']) or r['sit'] != 'Regular']
    cN = [
        {'k': 'data', 't': 'Data', 'w': 11, 'f': lambda r: r['data'], 'date': 1},
        {'k': 'es', 't': 'Entrada/Saída', 'w': 12, 'f': lambda r: r['es']},
        *([{'k': 'est', 't': 'Estabelecimento', 'w': 18, 'f': lambda r: r['filial']}] if multi else []),
        {'k': 'nro', 't': 'Nº NF', 'w': 9, 'f': lambda r: chave_nf(r) or r['nro'], 'z': '0'},
        {'k': 'ser', 't': 'Série', 'w': 6, 'f': lambda r: r['ser']},
        {'k': 'part', 't': 'Cliente/Fornecedor', 'w': 42, 'f': lambda r: r['part']},
        {'k': 'doc', 't': 'CNPJ/CPF', 'w': 19, 'f': lambda r: fmt_doc(r['doc'])},
        {'k': 'uf', 't': 'UF', 'w': 5, 'f': lambda r: r['uf']},
        {'k': 'cfop', 't': 'CFOP', 'w': 7, 'f': lambda r: r['cfop']},
        {'k': 'desc', 't': 'Descrição do CFOP', 'w': 55, 'f': lambda r: r['desc']},
        {'k': 'nat', 't': 'Natureza', 'w': 22, 'f': lambda r: r['nat']},
        {'k': 'vl', 't': 'Valor (R$)', 'w': 14, 'f': lambda r: r['vl'], 'num': 1},
        {'k': 'bc', 't': 'BC ICMS', 'w': 13, 'f': lambda r: r['bc'], 'num': 1},
        {'k': 'icms', 't': 'ICMS', 'w': 12, 'f': lambda r: r['icms'], 'num': 1},
        {'k': 'bcst', 't': 'BC ICMS ST', 'w': 13, 'f': lambda r: r['bcst'], 'num': 1},
        {'k': 'st', 't': 'ICMS ST', 'w': 12, 'f': lambda r: r['st'], 'num': 1},
        {'k': 'ipi', 't': 'IPI', 'w': 11, 'f': lambda r: r['ipi'], 'num': 1},
        *([{'k': 'sit', 't': 'Situação', 'w': 12, 'f': lambda r: r['sit']}] if any(r['sit'] != 'Regular' for r in notas) else []),
        {'k': 'chv', 't': 'Chave de acesso', 'w': 47, 'f': lambda r: r['chv']},
    ]
    LN, _ = tabela(wsN, 1, cN, notas)
    wsN.freeze_panes = 'A2'
    ultN = len(notas) + 1
    NT = "'Notas fiscais'!"
    rng = lambda k: f"{NT}${LN[k]}$2:${LN[k]}${ultN}"

    # ---- Produtos (itens das notas – C170)
    wsP['A1'] = 'Detalhamento dos produtos (itens das notas informados no SPED, registro C170)'
    wsP['A1'].font = F_SEC
    wsP['A2'] = ('O SPED traz os itens das notas recebidas de terceiros. Nas notas eletrônicas emitidas pela própria '
                 'empresa os itens não são informados no SPED.')
    wsP['A2'].font = F_NOTA
    if itens:
        inteiro = all(float(i['qtd']).is_integer() for i in itens)
        cP = [
            {'k': 'data', 't': 'Data', 'w': 11, 'f': lambda r: r['data'], 'date': 1},
            {'k': 'es', 't': 'Entrada/Saída', 'w': 12, 'f': lambda r: r['es']},
            *([{'k': 'est', 't': 'Estabelecimento', 'w': 18, 'f': lambda r: r['filial']}] if multi else []),
            {'k': 'nro', 't': 'Nº NF', 'w': 9, 'f': lambda r: chave_nf(r) or r['nro'], 'z': '0'},
            {'k': 'part', 't': 'Fornecedor/Cliente', 'w': 38, 'f': lambda r: r['part']},
            {'k': 'item', 't': 'Item', 'w': 5, 'f': lambda r: r['item'], 'z': '0'},
            {'k': 'cod', 't': 'Código', 'w': 12, 'f': lambda r: r['cod']},
            {'k': 'prod', 't': 'Descrição do produto', 'w': 56, 'f': lambda r: r['desc']},
            *([{'k': 'compl', 't': 'Complemento', 'w': 30, 'f': lambda r: r['compl']}] if any(i['compl'] for i in itens) else []),
            *([{'k': 'ncm', 't': 'NCM', 'w': 10, 'f': lambda r: r['ncm']}] if any(i['ncm'] for i in itens) else []),
            {'k': 'qtd', 't': 'Quantidade', 'w': 11, 'f': lambda r: r['qtd'], 'num': 1, 'z': '#,##0' if inteiro else '#,##0.000'},
            {'k': 'un', 't': 'Unidade', 'w': 8, 'f': lambda r: r['un']},
            {'k': 'unit', 't': 'Valor unitário (R$)', 'w': 15, 'f': lambda r: 0, 'num': 1, 'tot': False,
             'fx': lambda i, L: f'=IFERROR({L["vl"]}{i}/{L["qtd"]}{i},"")'},
            {'k': 'vl', 't': 'Valor do item (R$)', 'w': 15, 'f': lambda r: r['vl'], 'num': 1},
            {'k': 'dsc', 't': 'Desconto (R$)', 'w': 12, 'f': lambda r: r['dsc'], 'num': 1},
            {'k': 'cfop', 't': 'CFOP', 'w': 7, 'f': lambda r: r['cfop']},
            {'k': 'nat', 't': 'Natureza', 'w': 22, 'f': lambda r: r['nat']},
            {'k': 'bc', 't': 'BC ICMS', 'w': 12, 'f': lambda r: r['bc'], 'num': 1},
            {'k': 'icms', 't': 'ICMS', 'w': 11, 'f': lambda r: r['icms'], 'num': 1},
            {'k': 'bcst', 't': 'BC ICMS ST', 'w': 12, 'f': lambda r: r['bcst'], 'num': 1},
            {'k': 'st', 't': 'ICMS ST', 'w': 11, 'f': lambda r: r['st'], 'num': 1},
            {'k': 'ipi', 't': 'IPI', 'w': 10, 'f': lambda r: r['ipi'], 'num': 1},
        ]
        tabela(wsP, 4, cP, itens)
        wsP.freeze_panes = 'A5'
    else:
        wsP['A4'] = 'Nenhum item de nota (C170) foi informado nos SPEDs carregados.'

    # ---- Clientes e fornecedores (valores por fórmula sobre a aba Notas fiscais)
    grp = defaultdict(list)
    for r in notas:
        if nz(r['vl']):
            grp[r['es'], r['part'], r['doc']].append(r)
    parts = sorted(grp.items(), key=lambda kv: (kv[0][0], -sum(r['vl'] for r in kv[1])))
    dados_c = [{'es': es, 'part': pt, 'doc': dc, 'uf': g[0]['uf'], 'q': sum(r['first'] for r in g),
                'vv': sum(r['vl'] for r in g if r['nat'] == 'Venda'),
                'cc': sum(r['vl'] for r in g if r['nat'] == 'Compra para revenda'),
                'v': sum(r['vl'] for r in g)} for (es, pt, dc), g in parts]
    crit = lambda i, L: f'{rng("es")},$A{i},{rng("part")},$B{i},{rng("doc")},$C{i}'
    cC = [
        {'k': 'es', 't': 'Entrada/Saída', 'w': 12, 'f': lambda r: r['es']},
        {'k': 'part', 't': 'Cliente/Fornecedor', 'w': 44, 'f': lambda r: r['part']},
        {'k': 'doc', 't': 'CNPJ/CPF', 'w': 19, 'f': lambda r: fmt_doc(r['doc'])},
        {'k': 'uf', 't': 'UF', 'w': 5, 'f': lambda r: r['uf']},
        {'k': 'q', 't': 'Notas', 'w': 8, 'f': lambda r: r['q'], 'z': '0'},
        *([{'k': 'vv', 't': 'Vendas (R$)', 'w': 15, 'f': lambda r: 0, 'num': 1,
            'fx': lambda i, L: f'=SUMIFS({rng("vl")},{crit(i, L)},{rng("nat")},"Venda")'}] if any(nz(d['vv']) for d in dados_c) else []),
        *([{'k': 'cc', 't': 'Compras p/ revenda (R$)', 'w': 18, 'f': lambda r: 0, 'num': 1,
            'fx': lambda i, L: f'=SUMIFS({rng("vl")},{crit(i, L)},{rng("nat")},"Compra para revenda")'}] if any(nz(d['cc']) for d in dados_c) else []),
        {'k': 'v', 't': 'Valor total (R$)', 'w': 15, 'f': lambda r: 0, 'num': 1,
         'fx': lambda i, L: f'=SUMIFS({rng("vl")},{crit(i, L)})'},
    ]
    tabela(wsC, 1, cC, dados_c)
    wsC.freeze_panes = 'A2'

    # ---- Resumo: dados da empresa e valores consolidados (só o que tem valor), com visual formal
    R = wsR
    NC = 6
    widths(R, [44, 22, 58, 24, 16, 18])
    FILL_NAVY = PatternFill('solid', fgColor=NAVY)
    FILL_SEC = PatternFill('solid', fgColor='D9E3EE')
    FILL_ZEBRA = PatternFill('solid', fgColor='F5F8FB')
    FILL_LBL = PatternFill('solid', fgColor='EEF2F6')
    GRID = Side(style='thin', color='C9D3DD')
    B_GRID = Border(left=GRID, right=GRID, top=GRID, bottom=GRID)
    B_TOT = Border(left=GRID, right=GRID, top=Side(style='thin', color=NAVY), bottom=Side(style='double', color=NAVY))
    AL_R = Alignment(horizontal='right', vertical='center')
    AL_L = Alignment(horizontal='left', vertical='center', wrap_text=True, indent=1)

    def faixa(lin, texto, altura=20, fonte=None, fill=None, c2=NC):
        R.merge_cells(start_row=lin, start_column=1, end_row=lin, end_column=c2)
        c = R.cell(lin, 1, texto)
        c.font = fonte or Font(name='Arial', size=11, bold=True, color='FFFFFF')
        c.fill = fill or FILL_NAVY
        c.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        for j in range(2, c2 + 1):
            R.cell(lin, j).fill = fill or FILL_NAVY
        R.row_dimensions[lin].height = altura

    def celula(lin, col, v, fmt=None, bold=False, fill=None, borda=B_GRID, al=None, cor=None):
        c = R.cell(lin, col, v)
        c.font = Font(name='Arial', size=10, bold=bold, color=cor)
        c.border = borda
        if fill:
            c.fill = fill
        if fmt:
            c.number_format = fmt
        c.alignment = al or (AL_R if fmt else AL_L)
        return c

    def cabecalho(lin, titulos, c1=1):
        for j, t in enumerate(titulos, c1):
            direita = t.endswith('(R$)') or t in ('Notas', 'Notas no período', 'Valor')
            c = R.cell(lin, j, t)
            c.font = Font(name='Arial', size=9, bold=True, color=NAVY)
            c.fill = FILL_SEC
            c.border = B_GRID
            c.alignment = Alignment(horizontal='right' if direita else 'left', vertical='center', wrap_text=True,
                                    indent=0 if direita else 1)
        R.row_dimensions[lin].height = 20

    faixa(1, 'RESUMO FISCAL DE COMPRAS E VENDAS', 30, Font(name='Arial', size=15, bold=True, color='FFFFFF'))
    faixa(2, f'{empresa}   ·   Período de {ini:%d/%m/%Y} a {fim:%d/%m/%Y}', 20,
          Font(name='Arial', size=10, bold=True, color=NAVY), FILL_SEC)
    lin = 4

    faixa(lin, 'DADOS DA EMPRESA')
    lin += 1
    for rot, val in (('Razão social', empresa), ('Período de apuração', f'{ini:%d/%m/%Y} a {fim:%d/%m/%Y}'),
                     ('Fonte das informações', 'SPED Fiscal – EFD ICMS/IPI')):
        celula(lin, 1, rot, bold=True, fill=FILL_LBL)
        R.merge_cells(start_row=lin, start_column=2, end_row=lin, end_column=NC)
        celula(lin, 2, val)
        for j in range(3, NC + 1):
            R.cell(lin, j).border = B_GRID
        lin += 1
    lin += 1
    cabecalho(lin, ['Estabelecimento', 'CNPJ', 'Inscrição estadual', 'UF', 'Notas no período'])
    for k, e in enumerate(ests):
        lin += 1
        n = sum(r['first'] for r in e['rows'])
        z = FILL_ZEBRA if k % 2 else None
        celula(lin, 1, e['label'], fill=z)
        celula(lin, 2, fmt_doc(e['cnpj']), fill=z)
        celula(lin, 3, e['ie'], fill=z)
        celula(lin, 4, e['uf'], fill=z)
        if n:
            celula(lin, 5, n, '#,##0', fill=z)
        else:
            celula(lin, 5, 'sem movimento', fill=z, al=AL_R, cor='8A96A3')
    lin += 2

    soma = lambda es, nat=None: sum(r['vl'] for r in notas if r['es'] == es and (nat is None or r['nat'] == nat))
    sf = lambda es, nat=None: (f'=SUMIFS({rng("vl")},{rng("es")},"{es}"' + (f',{rng("nat")},"{nat}"' if nat else '') + ')')
    faixa(lin, 'VALORES CONSOLIDADOS (R$)', c2=2)
    lin += 1
    ref = {}
    grupos = [
        ('Saídas', [('Vendas', 'Saída', 'Venda'), ('Venda de ativo', 'Saída', 'Venda de ativo'),
                    ('Devolução de compras', 'Saída', 'Devolução de compra'), ('Transferências enviadas', 'Saída', 'Transferência'),
                    ('Remessas / Retornos', 'Saída', 'Remessa/Retorno'), ('Bonificações / Brindes', 'Saída', 'Bonificação/Brinde'),
                    ('Serviços / Fretes', 'Saída', 'Serviços/Fretes'), ('Outras saídas', 'Saída', 'Outras'), ('Total de saídas', 'Saída', None)]),
        ('Entradas', [('Compras para revenda', 'Entrada', 'Compra para revenda'), ('Compras uso/consumo/ativo', 'Entrada', 'Compra uso/consumo/ativo'),
                      ('Devolução de vendas', 'Entrada', 'Devolução de venda'), ('Transferências recebidas', 'Entrada', 'Transferência'),
                      ('Remessas / Retornos', 'Entrada', 'Remessa/Retorno'), ('Bonificações / Brindes', 'Entrada', 'Bonificação/Brinde'),
                      ('Serviços / Fretes', 'Entrada', 'Serviços/Fretes'), ('Outras entradas', 'Entrada', 'Outras'), ('Total de entradas', 'Entrada', None)]),
    ]
    for sec, itens_sec in grupos:
        vis = [x for x in itens_sec if nz(soma(x[1], x[2]))]
        if not vis:
            continue
        cabecalho(lin, [sec, 'Valor'])
        lin += 1
        for k, (nome, es, nat) in enumerate(vis):
            tot = nat is None
            z = None if tot else (FILL_ZEBRA if k % 2 else None)
            celula(lin, 1, nome, bold=tot, fill=z, borda=B_TOT if tot else B_GRID)
            celula(lin, 2, sf(es, nat), MONEY, bold=tot, fill=z, borda=B_TOT if tot else B_GRID)
            ref[es, nat] = lin
            lin += 1
        lin += 1
    vd, dv = soma('Saída', 'Venda'), soma('Entrada', 'Devolução de venda')
    cp, dc = soma('Entrada', 'Compra para revenda'), soma('Saída', 'Devolução de compra')
    nf_venda = sum(r['first'] for r in notas if r['es'] == 'Saída' and r['nat'] == 'Venda')
    B = lambda k: f'B{ref[k]}' if k in ref else '0'
    ind = []
    if nz(vd) or nz(dv):
        ind.append(('Vendas líquidas (vendas – devoluções)', f'={B(("Saída", "Venda"))}-{B(("Entrada", "Devolução de venda"))}', MONEY))
    if nz(cp) or nz(dc):
        ind.append(('Compras líquidas (compras – devoluções)', f'={B(("Entrada", "Compra para revenda"))}-{B(("Saída", "Devolução de compra"))}', MONEY))
    if len(ind) == 2 and nz((vd - dv) - (cp - dc)):
        ind.append(('Vendas líquidas – compras líquidas', 'DIF', MONEY))
    if nz(vd) and nf_venda:
        ind.append(('Ticket médio de venda (por nota)', f'={B(("Saída", "Venda"))}/{nf_venda}', MONEY))
    for nome, v in (('Notas fiscais de saída', sum(r['first'] for r in rows if r['es'] == 'Saída')),
                    ('Notas fiscais de entrada', sum(r['first'] for r in rows if r['es'] == 'Entrada'))):
        if v:
            ind.append((nome, v, '#,##0'))
    if ind:
        cabecalho(lin, ['Indicadores', 'Valor'])
        lin += 1
        base = lin
        for k, (nome, f, fmt) in enumerate(ind):
            z = FILL_ZEBRA if k % 2 else None
            celula(lin, 1, nome, fill=z)
            celula(lin, 2, f'=B{base}-B{base + 1}' if f == 'DIF' else f, fmt, fill=z)
            lin += 1
        lin += 1

    # Operações por CFOP
    cf = defaultdict(list)
    for r in notas:
        if r['cfop'] and nz(r['vl']):
            cf[r['es'], r['cfop']].append(r)
    cf = sorted(cf.items(), key=lambda kv: (kv[0][0], -sum(r['vl'] for r in kv[1])))
    faixa(lin, 'OPERAÇÕES POR CFOP')
    lin += 1
    cabecalho(lin, ['Entrada/Saída', 'CFOP', 'Descrição da operação', 'Natureza', 'Notas', 'Valor (R$)'])
    for k, ((es, c), g) in enumerate(cf):
        lin += 1
        z = FILL_ZEBRA if k % 2 else None
        celula(lin, 1, es, fill=z)
        celula(lin, 2, c, '0', fill=z, al=Alignment(horizontal='center', vertical='center'))
        celula(lin, 3, g[0]['desc'], fill=z)
        celula(lin, 4, g[0]['nat'], fill=z)
        celula(lin, 5, sum(r['first'] for r in g), '#,##0', fill=z)
        celula(lin, 6, f'=SUMIFS({rng("vl")},{rng("es")},$A{lin},{rng("cfop")},$B{lin})', MONEY, fill=z)
        R.row_dimensions[lin].height = 14 * -(-len(g[0]['desc']) // 62) + 4
    lin += 2
    R.cell(lin, 1, 'Quantidade de notas contada no SPED; valores somados da aba "Notas fiscais". '
                   'Linhas e colunas sem valor não são exibidas.').font = F_NOTA
    R.print_area = f'A1:F{lin}'

    # ---- impressão: A4 paisagem, ajustado à largura, cabeçalho com o cliente e número de página
    for ws, titulos in ((wsR, None), (wsN, '1:1'), (wsP, '4:4'), (wsC, '1:1')):
        ws.page_setup.orientation = 'portrait' if ws is wsR else 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = ws.page_margins.bottom = 0.6
        if titulos:
            ws.print_title_rows = titulos
        ws.oddHeader.left.text = f'{empresa} – {per}'
        ws.oddHeader.left.size = 9
        ws.oddFooter.left.text = 'Fonte: SPED EFD ICMS/IPI'
        ws.oddFooter.left.size = 8
        ws.oddFooter.right.text = 'Página &P de &N'
        ws.oddFooter.right.size = 8
    wb.active = 0
    wb.save(saida)
    return rows, itens


# ---------------------------------------------------------------- main
def coletar(args):
    arquivos = []
    for a in args:
        p = Path(a)
        if p.is_dir():
            arquivos += sorted(p.glob('*.txt'))
        elif p.suffix.lower() == '.txt':
            arquivos.append(p)
    return arquivos


def main():
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('entradas', nargs='*', help='arquivos .txt do SPED ou pastas que os contenham')
    ap.add_argument('-o', '--saida', help='arquivo .xlsx de saída')
    a = ap.parse_args()
    entradas = a.entradas
    if not entradas:
        import tkinter as tk
        from tkinter import filedialog
        tk.Tk().withdraw()
        entradas = list(filedialog.askopenfilenames(title='Selecione os arquivos SPED Fiscal (.txt)',
                                                    filetypes=[('SPED Fiscal', '*.txt')]))
        if not entradas:
            return 1
    arquivos = coletar(entradas)
    ests, erros = {}, []
    for f in arquivos:
        try:
            e = parse_sped(ler_texto(f), f.name)
            ests[e['cnpj'], e['ini']] = e
        except Exception as ex:  # arquivo .txt que não é SPED
            erros.append(str(ex))
    if not ests:
        print('Nenhum SPED Fiscal válido encontrado.', *erros, sep='\n')
        return 1
    ests = sorted(ests.values(), key=lambda e: (e['cnpj'], e['ini']))
    saida = Path(a.saida) if a.saida else Path(arquivos[0]).parent / (
        f"Compras e Vendas SPED - {ests[0]['empresa']} - {ests[0]['ini']:%m.%Y}.xlsx")
    rows, itens = gerar(ests, saida)
    for e in ests:
        n = sum(r['first'] for r in e['rows'])
        print(f"  {e['label']:<22} CNPJ {e['cnpj']}  {e['ini']:%d/%m/%Y}-{e['fim']:%d/%m/%Y}  {n} documentos"
              + ('  (sem movimento)' if not n else ''))
    s = sum(r['vl'] for r in rows if r['es'] == 'Saída')
    en = sum(r['vl'] for r in rows if r['es'] == 'Entrada')
    fmt = lambda v: f'{v:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
    print(f'  Total saídas R$ {fmt(s)} | Total entradas R$ {fmt(en)} | {len(itens)} itens de produtos')
    for e in erros:
        print('  Ignorado:', e)
    print(f'Planilha gerada: {saida}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
